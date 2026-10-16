from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / 'docs' / 'Informe_PROMETEO_Taller_Investigacion_II.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.8)
sec.left_margin = sec.right_margin = Inches(.85)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Title'].font.size = Pt(23)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 2'].font.size = Pt(12)
for name in ['Heading 1', 'Heading 2']:
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.space_after = Pt(7)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    doc.add_heading(text, 1)

def page():
    doc.add_page_break()

def marker(number, text):
    para = p(f'[Espacio para figura {number}: {text}]')
    para.runs[0].italic = True
    para.paragraph_format.space_before = Pt(9)
    para.paragraph_format.space_after = Pt(9)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, width in zip(t.columns, widths):
        c.width = Inches(width)
    for c, txt in zip(t.rows[0].cells, headers):
        c.text = txt
    for row in rows:
        for c, txt in zip(t.add_row().cells, row):
            c.text = txt
    for i, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            pr.append(OxmlElement('w:tblHeader'))
        for c, width in zip(row.cells, widths):
            c.width = Inches(width)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'DCE6F1' if i == 0 else 'FFFFFF'); tcpr.append(sh)
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:'+side); e.set(qn('w:w'), '90'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            tcpr.append(margins)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:'+side); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
            tcpr.append(borders)
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    run.font.size = Pt(9)
                    run.bold = i == 0
    return t

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Página ').font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
doc.core_properties.title = 'Informe de pruebas del proyecto PROMETEO'
doc.core_properties.subject = 'Taller de Investigación II'
doc.core_properties.author = ''

p('Informe de pruebas del proyecto PROMETEO', 'Title')
p('Taller de Investigación II', 'Subtitle')
p('Sistema web de tamizaje psicológico para estudiantes del TESVB')
p('Estudiante: ____________________    Grupo: __________\nDocente: ______________________\nFecha de ejecución: 8 de octubre de 2026')
h('1 Presentación del avance')
p('Se verificaron los módulos de registro, autenticación y cuestionarios, además de controles de acceso y procesamiento de texto. La suite final obtuvo 125 pruebas aprobadas y 1272 aserciones. El avance corresponde a pruebas funcionales y de integración aislada; siguen pendientes la integración real con MariaDB y BETO y el cifrado de los datos sensibles identificados.')
h('2 Objetivo')
p('Verificar el funcionamiento de los módulos de PROMETEO y documentar los resultados y las limitaciones respecto de las actividades del cronograma: pruebas de módulos del 2 de octubre y pruebas de integración y seguridad del 9 de octubre.')
h('3 Alcance y metodología')
p('Se revisaron las pruebas existentes, se agregaron seis casos y se ejecutó PHPUnit mediante Laravel. Se emplearon cuentas y respuestas sintéticas, SQLite temporal en memoria y respuestas HTTP simuladas del servicio de IA. No se utilizaron expedientes reales ni se enviaron textos a servicios externos.')
p('Entorno de ejecución: Windows, PHP 8.2.12, Laravel 12 y PHPUnit 11.5.55. La configuración de pruebas exige SQLite en memoria e impide solicitudes HTTP externas no simuladas.')
p('Aunque el cronograma denomina la primera actividad como pruebas unitarias, la evidencia principal corresponde a pruebas funcionales y de integración: intervienen rutas, permisos, modelos y base de datos. Estos resultados no constituyen una validación clínica de los instrumentos.')

page()
h('4 Matriz de pruebas')
p('La matriz resume los casos ejecutados y distingue los requisitos que aún no cuentan con evidencia de cumplimiento.')
table(['Caso', 'Verificación realizada', 'Resultado y alcance'], [
('MOD 01 Registro', 'Alta de usuario y persona, rol estudiante, consentimiento y rechazo de privilegios inyectados.', 'Aprobado en entorno aislado'),
('MOD 02 Autenticación', 'Acceso válido, contraseña incorrecta, cierre de sesión y bloqueo por intentos fallidos.', 'Aprobado en entorno aislado'),
('MOD 03 PHQ 9', 'Suma y persistencia de 9 respuestas; niveles y alertas para 0, 4, 5, 9, 10, 14, 15 y 27 puntos.', 'Aprobado según la lógica actual'),
('MOD 04 GAD 7', 'Suma y persistencia de 7 respuestas; niveles y alertas para 0, 4, 5, 9, 10, 14, 15 y 21 puntos.', 'Aprobado según la lógica actual'),
('MOD 05 DASS 21', '29 escenarios de dimensiones y límites; multiplicador por dos, severidad global, validación y transacciones.', 'Aprobado con SQLite'),
('INT 01 Análisis IA', 'Cola de procesamiento, estados, reintentos, contrato de respuesta y rechazo de resultados obsoletos.', 'Aprobado con IA simulada'),
('SEG 01 Roles', 'Rechazo de accesos no autorizados y permisos revocados; separación de información clínica.', 'Aprobado en casos cubiertos'),
('SEG 02 SQL', 'Dos expresiones SQL en autenticación y búsqueda de instrumento sin eludir controles.', 'Aprobado en dos entradas'),
('SEG 03 CSRF', 'Token ausente e incorrecto producen 419; token válido permite autenticación.', 'Aprobado en inicio de sesión'),
('SEG 04 Contraseñas', 'Hash almacenado verificable y contraseña omitida de la serialización.', 'Aprobado para contraseñas'),
('SEG 05 Cifrado', 'Revisión del diario, notas clínicas y respuestas de borradores.', 'Pendiente de corrección'),
('INT 02 Servicios reales', 'Flujo Laravel con MariaDB y BETO reales.', 'Pendiente de ejecución'),
], [1.25, 3.85, 1.7])

page()
h('5 Resultados de la ejecución')
table(['Ejecución', 'Pruebas aprobadas', 'Aserciones'], [('Suite previa', '119', '909'), ('Casos agregados', '6', '363'), ('Suite final conjunta', '125', '1272')], [3.1, 1.85, 1.85])
p('La suite final terminó sin fallos en 56.88 segundos. Una aserción comprueba una condición específica; su cantidad no equivale a un porcentaje de cobertura ni garantiza la ausencia de defectos.')
marker(1, 'resumen de la ejecución con 125 pruebas y 1272 aserciones')
h('6 Hallazgos y limitaciones')
doc.add_heading('Cifrado de datos sensibles', 2)
p('No se encontró cifrado de aplicación en los modelos revisados para el texto del diario, la impresión diagnóstica, la retroalimentación al estudiante y las respuestas de borradores. Las columnas son de texto o JSON. No se verificó el cifrado de discos, tablas ni respaldos. Por tanto, este requisito no se acredita como cumplido.')
p('El hash de contraseñas y el código anónimo del estudiante no sustituyen el cifrado clínico. La corrección requiere inventariar campos, gestionar claves, convertir registros existentes y verificar tanto el almacenamiento como la lectura autorizada.')
doc.add_heading('Integración y alcance de seguridad', 2)
p('Las respuestas simuladas y SQLite permiten comprobar la lógica, pero no la conexión con MariaDB, la configuración de BETO, los tiempos reales ni la concurrencia entre servicios. Las pruebas SQL y CSRF cubren puntos de entrada concretos; no constituyen una auditoría integral.')
doc.add_heading('Clasificación comprobada', 2)
p('PHQ-9 y GAD-7 usan actualmente cuatro etiquetas en el sistema: nulo hasta 4 puntos, leve de 5 a 9, moderado de 10 a 14 y severo desde 15. Las pruebas comprobaron esta implementación. Su correspondencia con el protocolo académico y clínico aprobado requiere una revisión independiente.')

page()
h('7 Actividades pendientes y cierre')
p('Para completar la integración real y la verificación del cifrado se propone el siguiente procedimiento en un entorno separado de pruebas.')
steps = [
'Preparar MariaDB, cola y servicio BETO de prueba. Registrar versiones y utilizar únicamente cuentas y textos sintéticos.',
'Aplicar las migraciones a una base exclusiva de pruebas. Mantener la protección de la suite actual que exige SQLite en memoria.',
'Registrar al estudiante, completar su expediente y aplicar cuestionarios con respuestas predeterminadas; contrastar los resultados con la matriz.',
'Enviar una entrada sintética al diario, ejecutar el trabajador de la cola y registrar las transiciones del análisis hasta su finalización.',
'Comprobar en MariaDB los campos persistidos frente a la respuesta real. Registrar cómo se envían PHQ-9 y GAD-7: el flujo documentado utiliza ceros de compatibilidad, que no representan cuestionarios aplicados.',
'Verificar vistas permitidas y denegadas con cuentas de estudiante, tutor y psicólogo. Probar fallo y recuperación de la IA solo en ese entorno separado.',
'Implementar el cifrado con conversión de registros previos; comprobar que el valor bruto no queda en claro y que la lectura autorizada funciona.',
'Registrar por caso fecha, versión, entrada, resultado esperado y observado, estado, evidencia e incidencia. Marcar como aprobado únicamente lo ejecutado y comprobado.',
]
for i, item in enumerate(steps, 1): p(f'{i}. {item}')
h('8 Conclusión')
p('Se verificaron mediante PHPUnit los módulos y controles incluidos en la matriz y se conservaron evidencias reproducibles. Se ampliaron los casos de límites de puntaje, entradas SQL, CSRF y contraseñas. La actividad de integración y seguridad permanece parcialmente completada hasta ejecutar el flujo real Laravel–MariaDB–BETO y resolver y comprobar el cifrado identificado como faltante.')

page()
h('9 Evidencia visual de los módulos')
p('Las siguientes figuras complementarán los resultados automatizados con capturas del sistema. Cada imagen debe identificar el caso y utilizar datos de prueba. Una captura de la interfaz ilustra el funcionamiento; por sí sola no acredita todos los controles automatizados.')
figures = [
(2, 'Registro de usuario', 'Formulario con datos sintéticos y confirmación de registro. Corresponde a MOD 01.'),
(3, 'Inicio de sesión', 'Acceso correcto y mensaje de rechazo ante contraseña incorrecta. Corresponde a MOD 02.'),
(4, 'Resultado de PHQ 9', 'Pantalla con nombre del instrumento, puntaje total y nivel mostrado. Corresponde a MOD 03.'),
(5, 'Resultado de GAD 7', 'Pantalla con nombre del instrumento, puntaje total y nivel mostrado. Corresponde a MOD 04.'),
(6, 'Resultado de DASS 21', 'Pantalla con puntajes y niveles de depresión, ansiedad y estrés. Corresponde a MOD 05.'),
(7, 'Acceso por rol', 'Mensaje de acceso denegado a una ruta restringida, indicando el rol utilizado. Corresponde a SEG 01.'),
]
for n, title, desc in figures:
    doc.add_heading(f'Figura {n} {title}', 2)
    p(desc)
    marker(n, 'captura por incorporar')
p('Evidencia adicional para INT 02: cuando se ejecute el flujo real, incorporar el análisis completado del diario y su persistencia en MariaDB. Esa evidencia aún no se presenta como obtenida.')

page()
h('10 Anexo de trazabilidad técnica')
p('Referencia del código antes de agregar pruebas:')
p('53b5acddacfa5b3232e25e89d43dea7f51828896')
table(['Casos', 'Clases de prueba'], [
('MOD 01 y MOD 02', 'RegistrationTest, AuthenticationTest y ConsentAndAccessTest'),
('MOD 03 y MOD 04', 'ResearchScheduleTest y QuestionnaireDraftTest'),
('MOD 05', 'ResearchScheduleTest y Dass21IntegrationTest'),
('INT 01', 'AnalisisNlpProcessingTest'),
('SEG 01', 'ConsentAndAccessTest, Dass21IntegrationTest y PendingExpedienteAuthorizationTest'),
('SEG 02 a SEG 04', 'ResearchScheduleTest'),
], [1.6, 5.2])
p('Las clases se encuentran en tests/Feature; las de autenticación están en tests/Feature/Auth. ResearchScheduleTest contiene los seis casos añadidos para esta actividad.')
doc.add_heading('Archivos de evidencia', 2)
p('Carpeta: docs/evidencias/taller-investigacion-ii')
for item in ['ejecucion.txt: salida de la suite final.', 'resultados-phpunit.xml: reporte estructurado de PHPUnit.', 'tests/Feature/ResearchScheduleTest.php: casos agregados, desde la raíz del proyecto.']:
    p(item)
doc.add_heading('Reproducción', 2)
p('Desde la carpeta del proyecto, con las dependencias instaladas y sin configuración en caché:')
for cmd in ['php artisan test --compact --filter=ResearchScheduleTest', 'php artisan test --compact', 'php artisan test --compact\n  --log-junit=docs/evidencias/taller-investigacion-ii/resultados-phpunit.xml']:
    para = p(cmd)
    for r in para.runs: r.font.name = 'Consolas'; r.font.size = Pt(9)
p('La última instrucción se ejecuta en una sola línea. El reporte XML se reemplaza en cada ejecución; conservar una copia fechada para comparar entregas.')
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
