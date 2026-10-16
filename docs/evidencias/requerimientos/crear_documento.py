from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[2] / 'Requerimientos_PROMETEO.docx'
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.7)
s.left_margin = s.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = d.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
d.styles['Normal'].font.size = Pt(11)
d.styles['Normal'].paragraph_format.space_after = Pt(6)
d.styles['Normal'].paragraph_format.line_spacing = 1.06
d.styles['Title'].font.size = Pt(26)
d.styles['Heading 1'].font.size = Pt(18)
d.styles['Heading 2'].font.size = Pt(12)
d.styles['Heading 2'].paragraph_format.space_before = Pt(10)
d.styles['Heading 2'].paragraph_format.space_after = Pt(4)
footer = s.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('PROMETEO  |  ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
for r in footer.runs: r.font.size = Pt(9)
d.core_properties.title = 'Requerimientos funcionales y no funcionales de PROMETEO'
d.core_properties.subject = 'Especificación del sistema de seguimiento estudiantil'
d.core_properties.author = 'Equipo del proyecto PROMETEO'

def p(t): return d.add_paragraph(t)
def h(t): d.add_heading(t, 1)
def page(t): d.add_page_break(); h(t)
def req(code, title, actor, body, acceptance, priority='Alta', state='Base funcional'):
    d.add_heading(code.replace('-', ' ') + '  ' + title, 2)
    a=p(f'Actor: {actor} · Prioridad: {priority} · Clasificación: {state}')
    a.paragraph_format.space_after=Pt(3)
    for r in a.runs: r.font.size=Pt(9)
    p(body)
    a=p('Aceptación: ' + acceptance)
    a.paragraph_format.space_after=Pt(8)

d.add_heading('Requerimientos funcionales y no funcionales', 0)
d.add_paragraph('PROMETEO', 'Subtitle')
p('Sistema web de seguimiento estudiantil')
p('Versión 1.0 · 9 de octubre de 2026')
h('1 Propósito y alcance')
p('Esta especificación define el comportamiento esperado de PROMETEO y los criterios para aceptar sus funciones y su calidad. Está dirigida al equipo de desarrollo, responsables de la institución, Control Escolar, tutoría y psicología. Sirve como base para acordar el alcance y organizar la validación del sistema.')
p('PROMETEO debe apoyar el registro de estudiantes, la organización académica, la aplicación de cuestionarios, el diario emocional, la identificación de señales de atención y el seguimiento profesional. Incluye administración de accesos, evolución de resultados, notificaciones internas y reportes de participación.')
p('El alcance no incluye diagnóstico automático definitivo, atención de emergencias, prescripción, mensajería externa de casos ni entrenamiento de modelos de IA. La interpretación profesional y las decisiones de atención corresponden a psicología.')
h('2 Actores del sistema')
for a,b in [
('Estudiante','registra sus datos, responde cuestionarios, escribe su diario y consulta su evolución y retroalimentación compartida.'),
('Tutor','gestiona estudiantes dentro de sus grupos autorizados y consulta información general de seguimiento.'),
('Psicólogo','consulta resultados autorizados, registra valoraciones y administra casos de atención.'),
('Control Escolar','administra catálogos académicos, estudiantes, tutores y asignaciones según sus permisos.'),
('Administrador','administra usuarios, roles, permisos, expedientes y catálogos; consulta reportes autorizados.')]:
    p(a + ': ' + b)
h('3 Convenciones de la especificación')
p('RF identifica requerimientos funcionales y RNF identifica requerimientos no funcionales. Prioridad alta indica funciones esenciales para operar o proteger información; media indica funciones de apoyo. Estas prioridades se proponen para la planificación.')
p('Base funcional significa que el comportamiento está descrito en el código o la documentación del proyecto; no certifica su despliegue. Base técnica identifica controles existentes cuya efectividad debe verificarse. Meta propuesta indica una condición de calidad que la institución debe aprobar y medir antes de asumirla como compromiso.')

page('4 Acceso y administración de usuarios')
req('RF-01','Registro de estudiantes','Visitante', 'El sistema deberá permitir crear una cuenta estudiantil con correo único, contraseña confirmada y datos personales obligatorios, sin permitir que el solicitante se otorgue privilegios administrativos.', 'Un registro válido crea la cuenta correspondiente; un correo repetido o un dato inválido se rechaza y señala el campo que debe corregirse.')
req('RF-02','Inicio y cierre de sesión','Todos los usuarios', 'El sistema deberá autenticar mediante correo y contraseña, dirigir al panel correspondiente al rol y permitir cerrar la sesión.', 'Credenciales válidas permiten acceder al panel autorizado; las inválidas no crean sesión. Después de salir, una ruta privada vuelve a exigir autenticación.')
req('RF-03','Recuperación de contraseña','Usuario registrado', 'El sistema deberá ofrecer recuperación de contraseña mediante un enlace enviado al correo registrado y permitir establecer una nueva contraseña válida.', 'En un entorno con correo configurado, el enlace válido permite actualizar la contraseña; un token inválido o vencido se rechaza.')
req('RF-04','Consentimiento de uso','Todos los usuarios', 'El sistema deberá mostrar el aviso de privacidad y exigir aceptación del consentimiento antes de permitir el uso de las funciones protegidas, conservando la fecha de aceptación.', 'Aceptar habilita el acceso y registra la fecha. Rechazar cierra la sesión. Intentar abrir directamente una función protegida sin aceptación no permite usarla.')
req('RF-05','Usuarios roles y permisos','Administrador', 'El sistema deberá permitir administrar usuarios, roles y permisos mediante las acciones autorizadas y aplicar esas restricciones en el servidor.', 'Un permiso otorgado habilita la acción correspondiente; al retirarlo, la misma solicitud directa es rechazada aunque el usuario conserve un enlace previo.')
req('RF-06','Expedientes y datos personales','Administrador y usuario autorizado', 'El sistema deberá vincular cuentas con personas y expedientes de estudiante, tutor o psicólogo; deberá identificar registros incompletos y permitir completarlos con autorización.', 'La cuenta sin expediente recibe un estado informativo y no puede ejecutar operaciones que requieren ese expediente. Completarlo habilita únicamente las acciones de su rol.')

page('5 Organización académica y evaluaciones')
req('RF-07','Catálogos académicos','Control Escolar y administrador', 'El sistema deberá permitir consultar y mantener carreras, ciclos escolares y grupos conforme a los permisos de cada actor y las restricciones de integridad.', 'Las altas y cambios válidos se conservan. Se rechazan identificadores inexistentes y operaciones sin permiso; cada catálogo muestra sus datos actualizados.')
req('RF-08','Gestión académica de estudiantes','Control Escolar y administrador', 'El sistema deberá permitir registrar y actualizar estudiantes, mantener su estado y consultar pendientes de asignación, incluyendo estudiantes sin grupo.', 'Un estudiante puede permanecer pendiente de grupo sin recibir una asignación ficticia; su expediente puede consultarse y completarse posteriormente.')
req('RF-09','Asignación e historial de grupos','Personal con permiso de asignación', 'El sistema deberá asignar, cambiar o retirar el grupo de un estudiante y guardar el movimiento junto con la actualización del expediente.', 'El historial registra el grupo de origen real. Repetir la misma operación no duplica movimientos y no se admiten nuevas asignaciones a grupos inactivos.')
req('RF-10','Tutores y grupos autorizados','Control Escolar y tutor', 'El sistema deberá gestionar tutores y sus asignaciones a grupos. El tutor deberá consultar y administrar estudiantes únicamente dentro del alcance autorizado.', 'Un tutor accede a sus grupos; una solicitud sobre un grupo o estudiante ajeno se rechaza. Los cambios de asignación se reflejan en su alcance de acceso.')
req('RF-11','Aplicación de cuestionarios','Estudiante', 'El sistema deberá ofrecer DASS-21, PHQ9 y GAD7, recuperar sus preguntas y validar respuestas completas con valores permitidos antes del envío.', 'DASS-21 exige 21 respuestas; PHQ9, 9; GAD7, 7. Respuestas faltantes o fuera de rango no generan una evaluación completada.')
req('RF-12','Borradores de cuestionarios','Estudiante', 'El sistema deberá guardar respuestas parciales por estudiante e instrumento, permitir guardado manual y automático y recuperar el avance al regresar.', 'El formulario recupera lo guardado y permite continuar. Un borrador no genera resultados, alertas, casos ni participación; un conflicto entre versiones pide recargar sin sobrescribir.')

page('6 Resultados y diario emocional')
req('RF-13','Cálculo y conservación de resultados','Sistema', 'El sistema deberá calcular los resultados según las reglas de cada instrumento y guardar respuestas, resultado, fecha y contexto académico disponible al responder.', 'Los casos de prueba de límites coinciden con las reglas implementadas. DASS-21 conserva tres dimensiones de 0 a 42; PHQ9 un total de 0 a 27 y GAD7 de 0 a 21. Un fallo revierte el envío completo.')
req('RF-14','Alertas por cuestionarios','Sistema y psicólogo', 'El sistema deberá generar alertas conforme a las reglas vigentes del proyecto: nivel moderado o severo en PHQ9 y GAD7; nivel máximo severo o extremadamente severo en DASS-21.', 'Una evaluación que cumple la regla genera una sola alerta vinculada a su origen y su caso de atención. Estas reglas deberán ser aprobadas por el responsable profesional antes de operar.')
req('RF-15','Historial y evolución','Estudiante y psicólogo autorizado', 'El sistema deberá mostrar el historial y la evolución temporal mediante series separadas para las tres dimensiones DASS-21, PHQ9 y GAD7, con filtro de periodo y tablas de apoyo.', 'El estudiante consulta solo su información; psicología requiere los permisos de detalle. Fechas y valores coinciden con los registros y no se mezclan escalas distintas.', priority='Media')
req('RF-16','Registro del diario','Estudiante', 'El sistema deberá guardar entradas propias de entre 20 y 3000 caracteres y permitir consultar su historial, sin exigir que termine el análisis de IA para confirmar el guardado.', 'Una entrada válida queda asociada al estudiante y su trabajo de análisis queda en cola. Longitudes inválidas se rechazan; una entrada guardada sigue disponible si la API no responde.')
req('RF-17','Análisis del diario y reintentos','Sistema y psicólogo autorizado', 'El sistema deberá procesar entradas en segundo plano, distinguir pendiente, procesando, completado y fallido y permitir reanálisis autorizado. Deberá conservar por separado clasificación, confianza, salida BETO y señal de atención.', 'Una respuesta válida actualiza los campos permitidos. Un fallo muestra un estado controlado y preserva el resultado anterior, si existe; una solicitud antigua no sobrescribe un reanálisis más reciente.')

page('7 Atención y seguimiento profesional')
req('RF-18','Apertura de casos','Sistema y psicólogo', 'El sistema deberá abrir un caso por evaluación con alerta, por diario cuya señal indique atención o por valoración de una evaluación sin alerta. El caso deberá conservar un vínculo con su origen.', 'Reprocesar un origen no duplica su caso. Dos evaluaciones o entradas diferentes del mismo estudiante pueden producir dos casos independientes.')
req('RF-19','Asignación y transferencia de casos','Psicólogo autorizado', 'El sistema deberá permitir asignar un caso a un profesional con expediente y permisos adecuados. Una vez asignado, solo su responsable podrá transferirlo o modificar su seguimiento.', 'La asignación exige motivo y registra responsable, autor y fecha. Otro profesional con acceso de lectura no puede modificarlo mientras tenga responsable.')
req('RF-20','Valoraciones y notas de seguimiento','Psicólogo responsable', 'El sistema deberá registrar la impresión privada y la retroalimentación compartida de una valoración, y mantener notas de seguimiento con autor y fecha. Una valoración deberá dejar el caso en seguimiento.', 'El estudiante solo consulta la retroalimentación compartida. Las notas previas no tienen una acción de edición o borrado; una corrección se registra como nota nueva.')
req('RF-21','Cierre y reapertura','Psicólogo responsable', 'El sistema deberá manejar los estados pendiente, asignado, seguimiento y cerrado. Cerrar requerirá seguimiento previo y motivo explícito; reabrir requerirá motivo.', 'Abrir la pantalla no cambia el estado. Una valoración no cierra automáticamente. Cerrar o reabrir registra el evento y no elimina las señales originales del análisis.')
req('RF-22','Consulta general para tutoría','Tutor', 'El sistema deberá mostrar el estado general de seguimiento de estudiantes activos en grupos autorizados, incluyendo personas con al menos un caso abierto.', 'Una persona con varios casos abiertos se cuenta una sola vez en el indicador de personas. Tutoría no recibe texto del diario, puntuaciones, origen del caso ni notas privadas.')

page('8 Avisos reportes y perfil')
req('RF-23','Notificaciones internas','Sistema y destinatario autorizado', 'El sistema deberá avisar de nuevos casos a los psicólogos habilitados para el origen y de retroalimentación compartida al estudiante correspondiente; permitirá filtrar y marcar avisos leídos o sin leer.', 'Reanalizar un caso existente no duplica el aviso de apertura. Abrir una notificación vuelve a comprobar el acceso; el aviso no almacena contenido clínico ni texto del diario.')
req('RF-24','Reportes de participación','Administrador psicólogo y Control Escolar con permiso', 'El sistema deberá presentar aplicaciones y participantes por instrumento y contexto académico, con filtros de fechas, carrera, grupo, ciclo y condición de asignación.', 'Se utilizan las fechas de aplicación y el contexto registrado entonces. Se distinguen sin grupo al responder e histórico sin contexto; participantes significa personas distintas dentro de cada fila.', priority='Media')
req('RF-25','Exportación de reportes','Usuario autorizado para reportes', 'El sistema deberá exportar los resultados agregados a CSV conservando los filtros y permisos de la consulta.', 'Pantalla y archivo contienen los mismos agregados. El CSV no incluye identificadores personales, puntuaciones, diarios ni notas; neutraliza valores que una hoja de cálculo interpretaría como fórmulas.', priority='Media')
req('RF-26','Edición de perfil','Usuario autenticado', 'El sistema deberá permitir actualizar los datos propios habilitados y cambiar la contraseña previa validación de las condiciones de seguridad.', 'Los cambios válidos persisten; intentar modificar el perfil de otra cuenta se rechaza y un dato inválido muestra un mensaje comprensible.', priority='Media')
req('RF-27','Fotografía de perfil','Usuario con datos personales completos', 'El sistema deberá aceptar JPEG, PNG o WebP de hasta 2 MB y 3000 por 3000 píxeles, generar una imagen privada de 512 por 512 y permitir reemplazarla o eliminarla.', 'El archivo generado no conserva metadatos y solo se sirve al propietario autenticado. Sin foto se muestran iniciales; formatos o tamaños inválidos se rechazan.', priority='Media')
req('RF-28','Preferencias de apariencia','Usuario autenticado', 'El sistema deberá guardar preferencias de tema, acento, densidad y movimiento y aplicarlas de manera consistente, conservando los colores semánticos de los estados.', 'Tras guardar y volver a ingresar se conservan las preferencias. La opción de movimiento reducido elimina animaciones decorativas sin ocultar gráficas informativas.', priority='Media')

page('9 Seguridad integridad y resiliencia')
def nf(code,title,body,acceptance, state='Base técnica',priority='Alta'):
    req(code,title,'Sistema y equipo de operación',body,acceptance,priority,state)
nf('RNF-01','Autorización y confidencialidad','Toda consulta o modificación protegida deberá verificar sesión, consentimiento, rol, permiso y propiedad o alcance del registro cuando corresponda.', 'Las pruebas de acceso directo, identificadores ajenos y permisos revocados deniegan todos los escenarios no autorizados, incluidos avisos, fotos y exportaciones.')
nf('RNF-02','Protección de sesiones','La operación productiva deberá utilizar HTTPS, cookies Secure y HttpOnly, protección CSRF y restricción de intentos de autenticación.', 'La configuración productiva y las solicitudes se verifican; cinco errores de acceso por combinación de correo e IP activan un bloqueo temporal de 60 segundos. Cerrar sesión invalida el acceso.')
nf('RNF-03','Minimización de información','Los registros de errores y trabajos de cola deberán evitar diarios, respuestas, nombres y cuerpos de respuesta de IA. Los errores al usuario no deberán mostrar credenciales ni detalles internos.', 'Con datos sintéticos, los fallos solo registran identificadores técnicos y códigos controlados; los trabajos de análisis contienen identificador y UUID, sin el texto del diario.')
nf('RNF-04','Integridad transaccional','Evaluaciones, resultados, borradores, movimientos académicos y cambios de caso deberán conservar consistencia mediante operaciones atómicas y referencias válidas.', 'Un fallo inducido durante una escritura revierte los cambios relacionados; no deja resultados parciales, movimientos aislados ni estados de caso sin su evento correspondiente.')
nf('RNF-05','Control de concurrencia','El sistema deberá impedir duplicación o sobrescritura causada por reenvíos, pestañas desactualizadas o procesos simultáneos en casos, borradores y análisis.', 'Dos solicitudes concurrentes sobre la misma versión producen una actualización válida y un rechazo controlado; un trabajo con UUID anterior no reemplaza el resultado vigente.')
nf('RNF-06','Tolerancia a fallos de IA','Los errores transitorios de conexión, HTTP 408, 429 y 5xx deberán tener como máximo tres intentos totales; contratos inválidos deberán fallar sin reintento automático.', 'Se comprueban esperas de 30 y 120 segundos, límite HTTP de 60 segundos y estados finales controlados. Caer la API no impide guardar nuevas entradas de diario.')

page('10 Metas de calidad y operación')
nf('RNF-07','Rendimiento','Las consultas ordinarias deberán responder en un máximo de 2 segundos y los guardados ordinarios en 3 segundos para al menos el 95 % de las solicitudes, con 50 usuarios concurrentes.', 'Una prueba de 15 minutos con datos sintéticos representativos mide percentil 95 y errores menores al 1 %. Se excluyen inferencia y exportaciones; se documentan equipo, volumen y mezcla de operaciones.', 'Meta propuesta')
nf('RNF-08','Disponibilidad','La aplicación web deberá alcanzar una disponibilidad mensual de al menos 99.5 %, excluyendo mantenimiento programado y comunicado.', 'Un monitor externo registra comprobaciones cada minuto y calcula tiempo disponible sobre tiempo comprometido. Se reporta por separado la disponibilidad del servicio de IA.', 'Meta propuesta')
nf('RNF-09','Respaldo y recuperación','La operación deberá realizar respaldos diarios protegidos y recuperar el servicio con pérdida máxima de 24 horas de información y en un plazo máximo de 4 horas.', 'Un ensayo trimestral restaura base, archivos y configuración en un entorno aislado, verifica registros y mide ambos límites de recuperación.', 'Meta propuesta')
nf('RNF-10','Usabilidad y accesibilidad','Los recorridos esenciales deberán operar con teclado, foco visible, etiquetas y errores comprensibles en español; las gráficas deberán ofrecer datos textuales y respetar movimiento reducido.', 'Se completan acceso, evaluación, diario y seguimiento sin ratón ni dependencia exclusiva del color, a 200 % de zoom y con lector de pantalla. Se documenta cada incidencia.', 'Meta propuesta', 'Media')
nf('RNF-11','Compatibilidad y adaptación','La interfaz deberá funcionar en las dos versiones estables más recientes de Chrome, Edge y Firefox disponibles al ejecutar la validación, en anchos de 360, 768 y 1366 píxeles.', 'Los recorridos esenciales no presentan controles inaccesibles ni pérdida de contenido. Las tablas amplias cuentan con desplazamiento contenido.', 'Meta propuesta', 'Media')
nf('RNF-12','Mantenibilidad y despliegue','Las entregas deberán conservar dependencias bloqueadas, migraciones versionadas, pruebas aisladas y una guía de actualización y reversión con respaldo.', 'Antes de liberar se ejecutan pruebas, compilación y revisión de configuración en un entorno aislado. Se valida el trabajador y no se ejecutan reinicializaciones destructivas sobre datos existentes.')

page('11 Trazabilidad y aceptación')
nf('RNF-13','Observabilidad operativa','El servicio deberá permitir detectar cola detenida, análisis fallidos y errores de infraestructura mediante estados y diagnósticos sin información sensible.', 'La caída simulada del trabajador deja entradas pendientes detectables. La comprobación operativa informa faltantes de configuración sin exponer secretos.')
nf('RNF-14','Retención y trazabilidad de información','La institución deberá definir plazos de conservación, responsables, eliminación y tratamiento de respaldos. La trazabilidad de consultas y eliminaciones deberá cubrirse conforme a esa política.', 'Se aprueba una política y se prueba su aplicación en datos sintéticos. Se documenta que el historial de casos existente no equivale a una auditoría inmutable y puede eliminarse en cascada con su origen.', 'Meta propuesta')
h('11 1 Fuentes internas de trazabilidad')
p('Las siguientes referencias permiten localizar el comportamiento y preparar la revisión técnica. Son rutas relativas a la raíz del proyecto.')
for text in [
'RF 01 a RF 06 y RF 26: routes/auth.php, routes/web.php, controladores Auth y Admin, solicitudes de registro y perfil, y pruebas Auth y PendingExpedienteAuthorizationTest.',
'RF 07 a RF 10: controladores ControlEscolar y Tutor, StudentGroupAssignmentService y ENTREGA-2-ESQUEMA.md en docs.',
'RF 11 a RF 15: EvaluacionController, Dass21ScoringService, Dass21EvaluationService, QuestionnaireDraftService y documentos DASS21-INTEGRACION y ENTREGA-6 en docs.',
'RF 16 a RF 22: DiarioController, ProcesarAnalisisNlp, CasoAtencionService y documentos ENTREGA-1-IA y ENTREGA-3-CASOS en docs.',
'RF 23 a RF 25: AvisoService, EvolucionController, ReporteController y ENTREGA-4-EVOLUCION-AVISOS-REPORTES.md en docs.',
'RF 27 y RF 28: ProfilePhotoController, ProfileController y ENTREGA-7-APARIENCIA-Y-PERFIL.md en docs. RNF: servicios y pruebas asociados, ENTREGA-5-OPERACION.md y metas de esta especificación.'
]:
    a=p(text)
    for r in a.runs: r.font.size=Pt(10)
h('11 2 Condiciones para aceptar la entrega')
p('Cada requerimiento deberá registrar resultado, evidencia, responsable y fecha de validación. Las funciones de prioridad alta deberán cumplir sus criterios; las metas propuestas deberán aprobarse antes de exigir su cumplimiento. La aceptación requiere revisión institucional del alcance y revisión profesional de las reglas de evaluación y atención.')
p('Antes de operar deberán quedar definidos el protocolo institucional ante señales de atención, los responsables y tiempos de respuesta, la conservación de datos y el contrato de IA para análisis solo de texto. Los valores PHQ9 y GAD7 enviados como cero por compatibilidad no deberán interpretarse como cuestionarios contestados.')
p('Responsable institucional: ____________________    Fecha: ______________\nResponsable técnico: ________________________    Fecha: ______________\nResponsable de psicología: ___________________    Fecha: ______________')
d.save(OUT)
print(OUT)
